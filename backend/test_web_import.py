"""网页导入（SSRF 防护、正文抽取）与 docx 解析回归测试。全部离线:
    cd backend && pytest test_web_import.py
"""
from __future__ import annotations

import socket

import httpx
import pytest
import web_import
from docx import Document
from file_service import FileService
from web_import import UnsafeUrlError, fetch_page, html_to_text

PUBLIC_IP = "93.184.216.34"


def _resolve_to(monkeypatch: pytest.MonkeyPatch, mapping: dict[str, str]) -> None:
    def fake_getaddrinfo(host, port, *args, **kwargs):
        if host not in mapping:
            raise socket.gaierror("unknown host")
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (mapping[host], port))]

    monkeypatch.setattr(web_import.socket, "getaddrinfo", fake_getaddrinfo)


def _client(handler) -> httpx.Client:
    return httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)


@pytest.mark.parametrize(
    "url",
    [
        "ftp://example.edu.cn/a",
        "file:///etc/passwd",
        "http://user:pw@example.edu.cn/",
        "http://example.edu.cn:6379/",
    ],
)
def test_rejects_bad_scheme_credentials_and_ports(url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    _resolve_to(monkeypatch, {"example.edu.cn": PUBLIC_IP})
    with pytest.raises(UnsafeUrlError):
        fetch_page(url, client=_client(lambda r: httpx.Response(200)))


@pytest.mark.parametrize("ip", ["127.0.0.1", "10.0.0.5", "192.168.1.1", "169.254.169.254", "::1"])
def test_rejects_private_addresses(ip: str, monkeypatch: pytest.MonkeyPatch) -> None:
    _resolve_to(monkeypatch, {"intranet.example": ip})
    with pytest.raises(UnsafeUrlError):
        fetch_page("http://intranet.example/", client=_client(lambda r: httpx.Response(200)))


def test_redirect_to_private_address_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    """公网地址 302 跳到元数据地址，第二跳也必须被拦下。"""
    _resolve_to(monkeypatch, {"news.example.edu.cn": PUBLIC_IP, "evil.internal": "169.254.169.254"})

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(302, headers={"location": "http://evil.internal/latest/meta-data"})

    with pytest.raises(UnsafeUrlError):
        fetch_page("https://news.example.edu.cn/a", client=_client(handler))


def test_fetches_article_text(monkeypatch: pytest.MonkeyPatch) -> None:
    _resolve_to(monkeypatch, {"news.example.edu.cn": PUBLIC_IP})
    html = (
        "<html><head><title>奖学金评审通知</title><script>var x=1</script></head>"
        "<body><nav>首页 | 新闻</nav><article><h1>通知</h1><p>申请截止 10 月 30 日。</p></article>"
        "<footer>版权所有</footer></body></html>"
    )

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, headers={"content-type": "text/html; charset=utf-8"}, text=html)

    page = fetch_page("https://news.example.edu.cn/a", client=_client(handler))

    assert page.title == "奖学金评审通知"
    assert "申请截止 10 月 30 日" in page.text
    assert "首页" not in page.text and "版权所有" not in page.text and "var x" not in page.text


def test_rejects_non_html(monkeypatch: pytest.MonkeyPatch) -> None:
    _resolve_to(monkeypatch, {"news.example.edu.cn": PUBLIC_IP})

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, headers={"content-type": "application/pdf"}, content=b"%PDF")

    with pytest.raises(ValueError, match="PDF"):
        fetch_page("https://news.example.edu.cn/a.pdf", client=_client(handler))


def test_rejects_oversized_body(monkeypatch: pytest.MonkeyPatch) -> None:
    _resolve_to(monkeypatch, {"news.example.edu.cn": PUBLIC_IP})
    monkeypatch.setattr(web_import, "MAX_BYTES", 10)

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, headers={"content-type": "text/plain"}, content=b"x" * 100)

    with pytest.raises(ValueError, match="上限"):
        fetch_page("https://news.example.edu.cn/a.txt", client=_client(handler))


def test_html_to_text_falls_back_to_body() -> None:
    title, text = html_to_text("<html><body><p>正文一</p><p>正文二</p></body></html>")
    assert title == ""
    assert text == "正文一\n正文二"


def test_docx_reads_paragraphs_and_tables(tmp_path) -> None:
    doc = Document()
    doc.add_paragraph("第一条 总则")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "国家奖学金"
    table.rows[0].cells[1].text = "10000 元"
    doc.add_paragraph("第二条 附则")
    path = tmp_path / "规章.docx"
    doc.save(path)

    text = FileService.read_file(str(path))

    assert text.splitlines() == ["第一条 总则", "国家奖学金 | 10000 元", "第二条 附则"]


def test_add_url_saves_markdown_under_tenant(tmp_path, monkeypatch: pytest.MonkeyPatch) -> None:
    import knowledge

    monkeypatch.setattr(knowledge, "_uploads_root", lambda: str(tmp_path))
    kb = knowledge.KnowledgeBase()
    kb._vector_db_failed = True  # 不连向量库

    page = web_import.WebPage(url="https://news.example.edu.cn/a", title="评审 通知", text="正文")
    info = kb.add_url("https://news.example.edu.cn/a", "tenantaaa111", fetch=lambda url: page)

    assert info.name == "评审_通知.md"
    assert info.source == "uploaded"
    saved = (tmp_path / "tenantaaa111" / "评审_通知.md").read_text(encoding="utf-8")
    assert "来源: https://news.example.edu.cn/a" in saved
